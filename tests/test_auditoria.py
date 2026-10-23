import io
from datetime import date
from decimal import Decimal
from unittest import mock

import openpyxl
import pandas as pd
import pytest

from auditor_cfdi import Configuracion, auditar_archivos, generar_excel
from auditor_cfdi.sat import EstadoSAT
from auditor_cfdi.tipo_cambio import TablaFIX

from .fabrica import cfdi40

ENERO = dict(fecha_inicio=date(2026, 1, 1), fecha_fin=date(2026, 1, 31))


def _uuid(n: int) -> str:
    return f"{n:08X}-0000-0000-0000-000000000000"


def _estados(**por_uuid):
    def consultor(c):
        return por_uuid.get(c.uuid, EstadoSAT("Vigente", "S - OK", validacion_efos="200"))
    return consultor


@pytest.fixture
def lote():
    return [
        ("mxn.xml", cfdi40(uuid=_uuid(1), subtotal="1000", descuento="100", total="1044")),
        ("usd.xml", cfdi40(uuid=_uuid(2), moneda="USD", tipo_cambio="17.5", subtotal="200",
                           total="232")),
        ("nota.xml", cfdi40(uuid=_uuid(3), tipo="E", subtotal="50", total="58")),
        ("pago.xml", cfdi40(uuid=_uuid(4), tipo="P", moneda="XXX", subtotal="0", total="0")),
        ("febrero.xml", cfdi40(uuid=_uuid(5), fecha="2026-02-01T00:00:00")),
        ("cancelada.xml", cfdi40(uuid=_uuid(6))),
        ("copia.xml", cfdi40(uuid=_uuid(1), subtotal="1000", descuento="100", total="1044")),
        ("roto.xml", b"<cfdi:Comprobante"),
    ]


def _fila(r, archivo):
    return r.detalle.set_index("Archivo").loc[archivo]


def test_reglas_de_inclusion(lote):
    consultor = _estados(**{_uuid(6): EstadoSAT("Cancelado", "S - OK")})
    r = auditar_archivos(lote, Configuracion(**ENERO, solo_periodo=True), consultor_sat=consultor)

    assert _fila(r, "mxn.xml")["Base MXN"] == 900.0
    assert _fila(r, "mxn.xml")["Contraprestación MXN"] == 45.0
    assert _fila(r, "usd.xml")["Base MXN"] == 3500.0           # 200 USD × 17.5 del CFDI
    assert _fila(r, "usd.xml")["Fuente TC"] == "CFDI"
    assert _fila(r, "nota.xml")["Base MXN"] == -50.0            # nota de crédito resta
    assert _fila(r, "nota.xml")["¿Incluida?"] == "Sí"

    motivos = {a: _fila(r, a)["Motivo de exclusión"] for a in
               ("pago.xml", "febrero.xml", "cancelada.xml", "copia.xml", "roto.xml")}
    assert "no forma parte de la base" in motivos["pago.xml"]
    assert motivos["febrero.xml"] == "Fuera del periodo auditado"
    assert motivos["cancelada.xml"] == "Cancelada ante el SAT"
    assert motivos["copia.xml"] == "UUID duplicado (ya incluido en mxn.xml)"
    assert motivos["roto.xml"].startswith("Error de lectura")

    assert r.base_mxn == pytest.approx(900 + 3500 - 50)
    assert r.contraprestacion == pytest.approx(45 + 175 - 2.5)
    assert r.por_periodo["Facturas"].tolist() == [3]
    assert set(r.exclusiones["Motivo"]) >= {"UUID duplicado", "Error de lectura"}


def test_solo_consulta_al_sat_lo_que_puede_contar(lote):
    consultados = []

    def consultor(c):
        consultados.append(c.uuid)
        return EstadoSAT("Vigente")

    auditar_archivos(lote, Configuracion(**ENERO, solo_periodo=True), consultor_sat=consultor)
    assert sorted(consultados) == [_uuid(1), _uuid(2), _uuid(3), _uuid(6)]

    consultados.clear()
    auditar_archivos(lote, Configuracion(**ENERO, calcular=False), consultor_sat=consultor)
    assert sorted(consultados) == [_uuid(i) for i in (1, 2, 3, 4, 5, 6)]  # sin duplicado ni roto


def test_por_defecto_suman_las_de_fuera_del_periodo(lote):
    r = auditar_archivos(lote, Configuracion(**ENERO), consultor_sat=_estados())
    febrero = _fila(r, "febrero.xml")
    assert febrero["¿Incluida?"] == "Sí" and febrero["En periodo"] == "No"
    assert febrero["Observaciones"] == "Emitida fuera del periodo"
    assert r.por_periodo["Periodo"].tolist() == ["2026-01", "2026-02"]


def test_modo_solo_verificar(lote):
    consultor = _estados(**{_uuid(6): EstadoSAT("Cancelado", "S - OK"),
                            _uuid(4): EstadoSAT("Error de conexión", "timeout")})
    r = auditar_archivos(lote, Configuracion(**ENERO, calcular=False), consultor_sat=consultor)
    assert r.contar_estado("Vigente") == 4 and r.contar_estado("Cancelado") == 1
    assert r.sin_verificar == 1 and r.errores_lectura == 1
    assert set(r.incidencias["Archivo"]) == {"cancelada.xml", "pago.xml", "copia.xml",
                                             "roto.xml", "febrero.xml"}
    assert "del cálculo" not in r.advertencias[0]


def test_opciones_de_configuracion(lote):
    config = Configuracion(**ENERO, solo_periodo=True, base="subtotal", restar_notas_credito=False,
                           fuente_tc="manual", tc_manual=Decimal("20"), tasa=Decimal("0.10"))
    r = auditar_archivos(lote, config, consultor_sat=_estados())
    assert _fila(r, "mxn.xml")["Base MXN"] == 1000.0
    assert _fila(r, "usd.xml")["Base MXN"] == 4000.0
    assert _fila(r, "nota.xml")["¿Incluida?"] == "No"
    assert r.contraprestacion == pytest.approx(100 + 400 + 100)  # mxn + usd + "cancelada" (aquí vigente)


def test_falla_del_sat_excluye_y_avisa(lote):
    consultor = _estados(**{_uuid(1): EstadoSAT("Error de conexión", "timeout")})
    r = auditar_archivos(lote, Configuracion(**ENERO), consultor_sat=consultor)
    assert _fila(r, "mxn.xml")["¿Incluida?"] == "No"
    assert "no pudieron verificarse" in r.advertencias[0]

    r = auditar_archivos(lote, Configuracion(**ENERO, incluir_no_verificadas=True),
                         consultor_sat=consultor)
    assert _fila(r, "mxn.xml")["¿Incluida?"] == "Sí"
    assert "sin verificar" in _fila(r, "mxn.xml")["Observaciones"]


def test_sin_verificar_sat():
    r = auditar_archivos([("a.xml", cfdi40())], Configuracion(**ENERO, verificar_sat=False))
    assert _fila(r, "a.xml")["¿Incluida?"] == "Sí"
    assert _fila(r, "a.xml")["Observaciones"] == "No verificada ante el SAT"


def test_efos_y_rfc_auditado():
    archivos = [("propia.xml", cfdi40(uuid=_uuid(1))),
                ("ajena.xml", cfdi40(uuid=_uuid(2), rfc_emisor="OTRO010101AAA"))]
    consultor = _estados(**{_uuid(1): EstadoSAT("Vigente", validacion_efos="100")})
    r = auditar_archivos(archivos, Configuracion(**ENERO, rfc_emisor="eku9003173c9"),
                         consultor_sat=consultor)
    assert _fila(r, "propia.xml")["Emisor en EFOS"] == "Sí"
    assert "EFOS" in _fila(r, "propia.xml")["Observaciones"]
    assert _fila(r, "ajena.xml")["Motivo de exclusión"] == "Emitida por un RFC distinto al auditado"


def test_tipo_cambio_banxico():
    archivos = [("usd.xml", cfdi40(moneda="USD", tipo_cambio="17", subtotal="100",
                                   fecha="2026-01-17T12:00:00"))]  # sábado
    tabla = TablaFIX({date(2026, 1, 15): Decimal("18.00"), date(2026, 1, 16): Decimal("18.50")})
    with mock.patch("auditor_cfdi.auditoria.descargar_fix", return_value=tabla) as fix:
        r = auditar_archivos(archivos, Configuracion(**ENERO, fuente_tc="banxico",
                                                     banxico_token="t"), consultor_sat=_estados())
    fix.assert_called_once()
    assert _fila(r, "usd.xml")["Tipo de cambio"] == 18.5
    assert _fila(r, "usd.xml")["Fuente TC"] == "Banxico FIX (DOF)"

    r = auditar_archivos(archivos, Configuracion(**ENERO, fuente_tc="banxico"),
                         consultor_sat=_estados())
    assert _fila(r, "usd.xml")["Tipo de cambio"] == 17.0
    assert "token" in r.advertencias[0]


def test_moneda_sin_tipo_de_cambio():
    r = auditar_archivos([("eur.xml", cfdi40(moneda="EUR"))], Configuracion(**ENERO),
                         consultor_sat=_estados())
    assert _fila(r, "eur.xml")["Motivo de exclusión"] == "Sin tipo de cambio para convertir a MXN"


def _hoja(r):
    ws = openpyxl.load_workbook(io.BytesIO(generar_excel(r))).active
    filas = [list(f) for f in ws.iter_rows(values_only=True)]
    encabezado = next(i for i, f in enumerate(filas) if "Estado SAT" in f)
    return ws, filas, encabezado


def test_excel_una_hoja(lote):
    consultor = _estados(**{_uuid(6): EstadoSAT("Cancelado", "S - OK")})
    r = auditar_archivos(lote, Configuracion(**ENERO), consultor_sat=consultor)
    ws, filas, enc = _hoja(r)
    assert ws.parent.sheetnames == ["Auditoría"]
    titulos = [t for t in filas[enc] if t]
    assert titulos[:2] == ["Estado SAT", "¿Cuenta?"] and "Contraprestación" in titulos
    assert len(filas) == enc + 1 + len(r.detalle) + 1  # encabezado + facturas + total
    j_liga = filas[enc].index("SAT") + 1
    assert ws.cell(enc + 2, j_liga).hyperlink.target.startswith("https://verificacfdi")
    indicadores = filas[enc - 5:enc]
    assert any(r.contraprestacion in f for f in indicadores)
    assert filas[-1][1] == "Total visible" and "SUBTOTAL(109" in str(filas[-1][15])
    notas = filas[enc].index("Notas")
    cancelada = next(f for f in filas if f[-1] == "cancelada.xml")
    assert cancelada[notas] == "Cancelada ante el SAT"


def test_excel_modo_verificacion(lote):
    r = auditar_archivos(lote, Configuracion(**ENERO, calcular=False), consultor_sat=_estados())
    _, filas, enc = _hoja(r)
    assert "Contraprestación" not in filas[enc] and "EFOS" in filas[enc]
    pago = next(f for f in filas if f[-1] == "pago.xml")
    assert pago[filas[enc].index("Notas")] in (None, "")  # en verificación no hay exclusiones


def test_excel_sin_facturas():
    r = auditar_archivos([], Configuracion(**ENERO))
    _, filas, enc = _hoja(r)
    assert filas[enc + 1][1] == "No se encontraron facturas en los archivos."


def test_modo_verificacion_tambien_suma(lote):
    consultor = _estados(**{_uuid(6): EstadoSAT("Cancelado", "S - OK")})
    r = auditar_archivos(lote, Configuracion(**ENERO, calcular=False), consultor_sat=consultor)
    assert _fila(r, "usd.xml")["Total MXN"] == 232 * 17.5
    assert _fila(r, "nota.xml")["Total MXN"] == -58.0
    assert pd.isna(_fila(r, "copia.xml")["Total MXN"])        # duplicado no suma
    # vigentes: mxn 1044 + usd 4060 + nota -58 + pago 0 + febrero 1160
    assert r.total_vigente_mxn == pytest.approx(1044 + 4060 - 58 + 0 + 1160)
    ws, filas, enc = _hoja(r)
    j = filas[enc].index("Total MXN")
    assert "SUBTOTAL(109" in str(filas[-1][j])
    assert any(r.total_vigente_mxn in f for f in filas[enc - 5:enc])


def test_conceptos_detallados():
    xml = cfdi40(subtotal="1000.00", conceptos=("Maniobras de carga", "Almacenaje"))
    r = auditar_archivos([("a.xml", xml)], Configuracion(**ENERO), consultor_sat=_estados())
    texto = _fila(r, "a.xml")["Conceptos"]
    assert texto.splitlines() == [
        "• Maniobras de carga — 2 Servicio × 250.00 = 500.00 MXN · clave 78101800",
        "• Almacenaje — 2 Servicio × 250.00 = 500.00 MXN · clave 78101800",
    ]
    assert r.conceptos["Descripción"].tolist() == ["Maniobras de carga", "Almacenaje"]
    assert r.conceptos["Importe neto MXN"].sum() == 1000.0
    ws, filas, enc = _hoja(r)
    j = filas[enc].index("Conceptos")
    assert filas[enc + 1][j] == texto
    assert ws.row_dimensions[enc + 2].height > 18  # fila más alta para los dos renglones
